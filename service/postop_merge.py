import argparse
import csv
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

CSV_ENCODING = "cp932"
DAYS_FROM = 21
DAYS_TO = 180
SURGERY_DATE_HEADER = "手術日"
PATIENT_ID_HEADER = "患者ID"
# 測定値が空欄かどうかの判定に含めない列
NON_MEASUREMENT_COLUMNS = frozenset({"日付", "ID", "SEQ", "機種", "測定日時"})
# CSV の「日付」列は 2 桁年（例: 24/03/23）
CSV_DATE_FORMAT = "%y/%m/%d"
DATETIME_FORMATS = ("%Y/%m/%d %H:%M", "%Y/%m/%d", CSV_DATE_FORMAT)

CsvRow = dict[str, str]


def load_csv(csv_path: Path) -> tuple[list[str], list[CsvRow]]:
    with csv_path.open(encoding=CSV_ENCODING, newline="") as file:
        reader = csv.DictReader(file)
        rows = list(reader)
    return list(reader.fieldnames or []), rows


def group_by_id(rows: list[CsvRow]) -> dict[str, list[CsvRow]]:
    grouped: dict[str, list[CsvRow]] = {}
    for row in rows:
        grouped.setdefault(row["ID"].strip(), []).append(row)
    return grouped


def parse_date(text: str) -> date:
    return datetime.strptime(text.strip(), CSV_DATE_FORMAT).date()


def is_blank_measurement(row: CsvRow) -> bool:
    return all(
        not (value or "").strip()
        for column, value in row.items()
        if column not in NON_MEASUREMENT_COLUMNS
    )


def select_closest_row(rows: list[CsvRow], surgery_date: date) -> CsvRow | None:
    """手術日の 21〜180 日後の行のうち、手術日に最も近い行を返す。

    同じ日に複数行ある場合は、測定値のある行を優先し SEQ の小さい行を採用する。
    """
    start = surgery_date + timedelta(days=DAYS_FROM)
    end = surgery_date + timedelta(days=DAYS_TO)
    candidates = [row for row in rows if start <= parse_date(row["日付"]) <= end]
    if not candidates:
        return None
    return min(
        candidates,
        key=lambda row: (
            parse_date(row["日付"]),
            is_blank_measurement(row),
            int(row.get("SEQ") or 0),
        ),
    )


def to_cell_value(text: str) -> int | float | datetime | str | None:
    """CSV の文字列を Excel で扱いやすい数値・日時に変換する。"""
    text = text.strip()
    if not text:
        return None
    for number_type in (int, float):
        try:
            return number_type(text)
        except ValueError:
            pass
    for datetime_format in DATETIME_FORMATS:
        try:
            return datetime.strptime(text, datetime_format)
        except ValueError:
            pass
    return text


def append_columns(
    sheet: Worksheet,
    prefix: str,
    fieldnames: list[str],
    rows_by_id: dict[str, list[CsvRow]],
) -> None:
    """シートの右端に CSV の全列を追加し、対象者ごとに選んだ 1 行を書き込む。"""
    header = [cell.value for cell in sheet[1]]
    date_column = header.index(SURGERY_DATE_HEADER) + 1
    id_column = header.index(PATIENT_ID_HEADER) + 1
    first_column = sheet.max_column + 1

    for offset, name in enumerate(fieldnames):
        sheet.cell(row=1, column=first_column + offset, value=f"{prefix}_{name}")

    for row_index in range(2, sheet.max_row + 1):
        surgery = sheet.cell(row=row_index, column=date_column).value
        patient_id = sheet.cell(row=row_index, column=id_column).value
        if not isinstance(surgery, datetime) or patient_id is None:
            continue
        selected = select_closest_row(
            rows_by_id.get(str(patient_id), []), surgery.date()
        )
        if selected is None:
            continue
        for offset, name in enumerate(fieldnames):
            sheet.cell(
                row=row_index,
                column=first_column + offset,
                value=to_cell_value(selected[name] or ""),
            )


def merge_postop_data(
    target_path: Path, vaiop_path: Path, refkeratometer_path: Path
) -> Path:
    workbook = load_workbook(target_path)
    sheet = workbook.active
    if not isinstance(sheet, Worksheet):
        raise ValueError(f"{target_path} にワークシートがありません")

    for prefix, csv_path in (("vaiop", vaiop_path), ("ref", refkeratometer_path)):
        fieldnames, rows = load_csv(csv_path)
        append_columns(sheet, prefix, fieldnames, group_by_id(rows))

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = target_path.with_name(f"{target_path.stem}_{timestamp}.xlsx")
    workbook.save(output_path)
    return output_path


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(
        description="targetdata.xlsx に術後 21〜180 日の vaiop・refkeratometer データを追加する"
    )
    parser.add_argument(
        "data_dir",
        nargs="?",
        type=Path,
        default=project_root / "data",
        help="targetdata.xlsx・vaiop.csv・refkeratometer.csv のあるディレクトリ",
    )
    args = parser.parse_args()

    output_path = merge_postop_data(
        args.data_dir / "targetdata.xlsx",
        args.data_dir / "vaiop.csv",
        args.data_dir / "refkeratometer.csv",
    )
    print(f"{output_path.resolve()} に出力しました")


if __name__ == "__main__":
    main()
