from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from service.postop_merge import merge_postop_data, select_closest_row, to_cell_value

SURGERY_DATE = date(2026, 1, 1)


def make_row(day: str, seq: str = "1", value: str = "1.0") -> dict[str, str]:
    return {"日付": day, "ID": "1", "SEQ": seq, "R_S": value}


def test_select_closest_row_boundaries() -> None:
    rows = [make_row("26/01/21"), make_row("26/01/22"), make_row("26/04/01")]
    assert select_closest_row(rows, SURGERY_DATE) == rows[1]
    assert select_closest_row([rows[0]], SURGERY_DATE) is None
    assert select_closest_row([rows[2]], SURGERY_DATE) == rows[2]
    assert select_closest_row([make_row("26/04/02")], SURGERY_DATE) is None


def test_select_closest_row_prefers_closest_to_28_days() -> None:
    day21 = make_row("26/01/22")
    day30 = make_row("26/01/31")
    assert select_closest_row([day21, day30], SURGERY_DATE) == day30


def test_select_closest_row_prefers_later_day_on_tie() -> None:
    day26 = make_row("26/01/27")
    day30 = make_row("26/01/31")
    assert select_closest_row([day26, day30], SURGERY_DATE) == day30


def test_select_closest_row_prefers_measured_row_then_smallest_seq() -> None:
    blank = make_row("26/02/01", seq="1", value="")
    later_seq = make_row("26/02/01", seq="3")
    earlier_seq = make_row("26/02/01", seq="2")
    assert select_closest_row([blank, later_seq, earlier_seq], SURGERY_DATE) == earlier_seq


def test_select_closest_row_adopts_blank_row_when_it_is_closest() -> None:
    blank = make_row("26/02/01", value="")
    measured = make_row("26/02/02")
    assert select_closest_row([measured, blank], SURGERY_DATE) == blank


def test_to_cell_value() -> None:
    assert to_cell_value("") is None
    assert to_cell_value("83") == 83
    assert to_cell_value("-0.75") == -0.75
    assert to_cell_value("2026/6/11 8:24") == datetime(2026, 6, 11, 8, 24)
    assert to_cell_value("24/03/23") == datetime(2024, 3, 23)
    assert to_cell_value("NIDEK ARK-1s") == "NIDEK ARK-1s"


def test_merge_postop_data(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["手術日", "患者ID"])
    sheet.append([datetime(2026, 1, 1), 1])
    sheet.append([datetime(2026, 1, 1), 2])
    target_path = tmp_path / "targetdata.xlsx"
    workbook.save(target_path)
    (tmp_path / "vaiop.csv").write_text(
        "日付,ID,304R\n26/02/01,1,13\n26/01/05,2,15\n", encoding="cp932"
    )
    (tmp_path / "refkeratometer.csv").write_text(
        "日付,ID,SEQ,R_S\n26/02/01,1,1,-1.25\n", encoding="cp932"
    )

    output_path = merge_postop_data(
        target_path, tmp_path / "vaiop.csv", tmp_path / "refkeratometer.csv"
    )

    assert output_path.name.startswith("targetdata_")
    output_sheet = load_workbook(output_path).active
    assert output_sheet is not None
    rows = list(output_sheet.iter_rows(values_only=True))
    assert rows[0] == (
        "手術日", "患者ID", "vaiop_日付", "vaiop_ID", "vaiop_304R",
        "ref_日付", "ref_ID", "ref_SEQ", "ref_R_S",
    )
    assert rows[1][2:] == (datetime(2026, 2, 1), 1, 13, datetime(2026, 2, 1), 1, 1, -1.25)
    assert rows[2][2:] == (None,) * 7
